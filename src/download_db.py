import firebase_admin
from firebase_admin import credentials, firestore
import openpyxl
from datetime import datetime
import shutil

CREDENTIAL_PATH = './firebase-key.json'
if not firebase_admin._apps:
    cred = credentials.Certificate(CREDENTIAL_PATH)
    firebase_admin.initialize_app(cred)
db = firestore.client()

print("\n1. 파이어베이스 현장 데이터 로딩 중...")
owners_ref = db.collection('owners').stream()
web_data = {doc.id: doc.to_dict() for doc in owners_ref}

EXCEL_PATH = r"C:\Users\이민후\OneDrive - 주식회사 광장건설\_주_광장건설 정비 PJ - 송파구 문정동 28-1번지 일원\문정동28-1 토지조서.xlsx"
current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
SAVE_PATH = EXCEL_PATH.replace(".xlsx", f"_현장반영_{current_time}.xlsx")

print("2. 원본 엑셀 파일 복사 및 로드 중...")
shutil.copy2(EXCEL_PATH, SAVE_PATH)
wb = openpyxl.load_workbook(SAVE_PATH)
ws = wb['통합']

# --- 개선된 열(Column) 탐색 로직 ---
# 엑셀의 셀 병합 때문에 글자가 숨어있는 것을 방지하기 위해 2~5번째 줄을 샅샅이 뒤집니다.
print("3. 엑셀의 타겟 열(Column) 위치 탐색 중...")
col_map = {}
for r in range(2, 6):
    for cell in ws[r]:
        if cell.value:
            clean_name = str(cell.value).replace('\n', '').replace(' ', '')
            col_map[clean_name] = cell.column

# 디버깅을 위해 파이썬이 엑셀에서 열을 제대로 찾았는지 터미널에 출력해 줍니다.
target_cols = ['연번', '조합설립동의서', '신분증', '개인정보', '성향', '상담내용']
print("   [열 인식 결과]")
for t_col in target_cols:
    if t_col in col_map:
        print(f"   - ✅ '{t_col}' 열 인식 성공 (엑셀 {col_map[t_col]}번째 칸)")
    else:
        print(f"   - ❌ '{t_col}' 열을 찾지 못했습니다! (엑셀 헤더 이름 확인 필요)")

update_count = 0

print("\n4. 데이터 엑셀 주입 시작...")
for row in range(5, ws.max_row + 1):
    id_cell_value = ws.cell(row=row, column=col_map.get('연번', 2)).value
    
    if id_cell_value is None or not str(id_cell_value).replace('.0', '').isdigit():
        continue
        
    doc_id = str(int(float(id_cell_value)))
    
    if doc_id in web_data:
        data = web_data[doc_id]
        wrote_any = False
        logs = []
        
        # 1) 조합설립동의서
        if '조합설립동의서' in col_map:
            if data.get('agreed', False):
                ws.cell(row=row, column=col_map['조합설립동의서']).value = 'O'
                logs.append('동의(O)')
                wrote_any = True
            else:
                ws.cell(row=row, column=col_map['조합설립동의서']).value = ''
                
        # 2) 신분증
        if '신분증' in col_map:
            if data.get('idCopy', False):
                ws.cell(row=row, column=col_map['신분증']).value = 'O'
                logs.append('신분증(O)')
                wrote_any = True
            else:
                ws.cell(row=row, column=col_map['신분증']).value = ''
                
        # 3) 개인정보
        if '개인정보' in col_map:
            if data.get('privacyConsent', False):
                ws.cell(row=row, column=col_map['개인정보']).value = 'O'
                logs.append('개인정보(O)')
                wrote_any = True
            else:
                ws.cell(row=row, column=col_map['개인정보']).value = ''
                
        # 4) 성향
        if '성향' in col_map and data.get('disposition'):
            ws.cell(row=row, column=col_map['성향']).value = data.get('disposition')
            logs.append('성향')
            wrote_any = True
            
        # 5) 상담내용
        if '상담내용' in col_map:
            memos = data.get('memoHistory', [])
            if memos:
                memo_texts = [f"[{m.get('date')}] {m.get('author')}: {m.get('text')}" for m in memos]
                ws.cell(row=row, column=col_map['상담내용']).value = '\n'.join(memo_texts)
                logs.append(f'메모({len(memos)}건)')
                wrote_any = True
            else:
                ws.cell(row=row, column=col_map['상담내용']).value = ''

        # 하나라도 O 표시나 메모가 적혔다면 화면에 출력
        if wrote_any:
            print(f"   -> [작성됨] 연번 {doc_id.zfill(3)} ({data.get('nm', '')}) : {', '.join(logs)}")
            update_count += 1

wb.save(SAVE_PATH)
print("-" * 40)
if update_count == 0:
    print("⚠️ 엑셀에 반영된 O 표시나 메모가 단 0건입니다.")
    print("   이유: 웹앱(파이어베이스)에 입력된 데이터가 없기 때문입니다.")
    print("   테스트를 위해 스마트폰이나 웹앱을 열고, 임의의 소유자에게 동의(O)와 메모를 추가한 뒤 다시 돌려보세요!")
else:
    print(f"✅ 다운로드 완료: 현장 데이터 {update_count}건이 엑셀에 반영되었습니다.")
print(f"📂 백업 및 반영된 파일 위치:\n   {SAVE_PATH}")