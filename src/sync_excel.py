"""Dashboard -> original workbook. Default is read-only; --apply saves a backup."""
import argparse
from collections import Counter
from datetime import datetime
import hashlib
import os
from pathlib import Path
import shutil
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import firebase_admin
from firebase_admin import credentials, firestore
import openpyxl
from openpyxl.cell.cell import MergedCell

TARGET = Path(r'C:\Users\이민후\OneDrive - 주식회사 광장건설\_주_광장건설 정비 PJ - 10.정비PJ 송파문정 28번지일대\토지조서 총괄.xlsx')
FIELDS = {'연락처': 'contact', '조합설립동의서': 'agreed', '신분증': 'idCopy',
          '개인정보': 'privacyConsent', '성향': 'disposition', '상담내용': 'memoHistory'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sync(apply=False):
    lock = TARGET.with_name('~$' + TARGET.name)
    if lock.exists():
        raise RuntimeError('Excel is open. Close it before retrying; original unchanged.')
    original_hash = digest(TARGET)
    if not firebase_admin._apps:
        firebase_admin.initialize_app(credentials.Certificate(str(Path(__file__).with_name('firebase-key.json'))))
    data = {d.id: d.to_dict() for d in firestore.client().collection('owners').stream(timeout=120)}
    if not data:
        raise RuntimeError('Empty dashboard response; original unchanged.')
    wb = openpyxl.load_workbook(TARGET)
    ws = wb['통합']
    columns = {}
    for row in ws.iter_rows(min_row=2, max_row=5):
        for cell in row:
            label = ''.join(str(cell.value or '').split())
            if label in {'연번', *FIELDS}:
                if label in columns and columns[label] != cell.column:
                    raise RuntimeError('Ambiguous column: ' + label)
                columns[label] = cell.column
    missing = {'연번', *FIELDS} - columns.keys()
    if missing:
        raise RuntimeError('Missing columns: ' + ', '.join(sorted(missing)))
    changes = Counter()
    matched = set()
    expected = {}
    for row in range(5, ws.max_row + 1):
        sn = ws.cell(row, columns['연번']).value
        try:
            number = float(sn)
            if not number.is_integer():
                continue
            key = str(int(number))
        except (ValueError, TypeError):
            continue
        if key not in data:
            continue
        matched.add(key)
        for header, field in FIELDS.items():
            if field not in data[key]:
                continue  # Missing fields must never erase existing Excel values.
            value = data[key][field]
            if field in ('agreed', 'idCopy', 'privacyConsent'):
                if not isinstance(value, bool):
                    raise RuntimeError('Invalid boolean field: ' + field)
                value = 'O' if value else None
            elif field == 'memoHistory':
                if not isinstance(value, list):
                    raise RuntimeError('Invalid memo history')
                value = '\n'.join(f"[{m.get('date', '')}] {m.get('author', '')}: {m.get('text', '')}" for m in value) or None
            else:
                if value is not None and not isinstance(value, str):
                    raise RuntimeError('Invalid text field: ' + field)
                value = value or None
            if isinstance(value, str) and len(value) > 32767:
                raise RuntimeError('Text exceeds Excel cell limit; original unchanged.')
            cell = ws.cell(row, columns[header])
            if isinstance(cell, MergedCell):
                raise RuntimeError('Merged target cell requires review: ' + cell.coordinate)
            if cell.data_type == 'f':
                raise RuntimeError('Target cell contains formula: ' + cell.coordinate)
            if (cell.value or None) != value:
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
                expected[cell.coordinate] = value
                changes[header] += 1
    if not matched:
        raise RuntimeError('No matching owner IDs; original unchanged.')
    print(f'Matched owners: {len(matched)}; changed cells: {sum(changes.values())}')
    print(dict(changes))
    if not apply or not changes:
        print('Read-only check completed.' if not apply else 'Already up to date.')
        return
    fd, name = tempfile.mkstemp(prefix='.dashboard-sync-', suffix='.xlsx', dir=TARGET.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        # Edit only target cell XML; preserve drawings, other sheets and all ZIP parts.
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        uri = ns['s']
        with zipfile.ZipFile(TARGET) as source:
            book = ET.fromstring(source.read('xl/workbook.xml'))
            sheet = next(s for s in book.find('s:sheets', ns) if s.attrib['name'] == '통합')
            rel_id = sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
            rels = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
            part = next(r.attrib['Target'] for r in rels if r.attrib['Id'] == rel_id)
            part = part.lstrip('/') if part.startswith('/') else 'xl/' + part
            xml = source.read(part).decode('utf-8')
            import re
            # Preserve namespaces and extensions verbatim; replace only known cell nodes.
            for coordinate, value in expected.items():
                pattern = r'<c\b(?=[^>]*\br="' + re.escape(coordinate) + r'")[^>]*?(?:/>|>.*?</c>)'
                match = re.search(pattern, xml, re.DOTALL)
                if not match:
                    raise RuntimeError('Target cell XML not found: ' + coordinate)
                cell_xml = ET.fromstring(match.group())
                attrs = dict(cell_xml.attrib)
                attrs.pop('t', None)
                from xml.sax.saxutils import quoteattr, escape
                if value is not None:
                    attrs['t'] = 'inlineStr'
                replacement = '<c ' + ' '.join(k + '=' + quoteattr(v) for k, v in attrs.items()) + '>'
                if value is not None:
                    replacement += '<is><t xml:space="preserve">' + escape(value) + '</t></is>'
                replacement += '</c>'
                xml = xml[:match.start()] + replacement + xml[match.end():]
            with zipfile.ZipFile(temporary, 'w', zipfile.ZIP_DEFLATED) as output:
                for item in source.infolist():
                    output.writestr(item, xml.encode('utf-8') if item.filename == part else source.read(item.filename))
        check = openpyxl.load_workbook(temporary)
        if check.sheetnames != wb.sheetnames or any(check['통합'][c].value != v for c, v in expected.items()):
            raise RuntimeError('Saved workbook verification failed.')
        check.close()
        wb.close()
        if lock.exists() or digest(TARGET) != original_hash:
            raise RuntimeError('Original changed or opened during sync; retry later.')
        backup_dir = TARGET.parent / '대시보드_자동반영_백업'
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f'{TARGET.stem}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx'
        shutil.copy2(TARGET, backup)
        if digest(backup) != original_hash or digest(TARGET) != original_hash or lock.exists():
            raise RuntimeError('Original changed during backup; original not replaced.')
        os.replace(temporary, TARGET)
        print('Updated: ' + str(TARGET))
        print('Backup: ' + str(backup))
    finally:
        temporary.unlink(missing_ok=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    sync(parser.parse_args().apply)
